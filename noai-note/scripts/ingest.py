#!/usr/bin/env python3
"""noai-note Minimal Ingester: Multi-source Noise Stripping & Meeting Chunking."""
import argparse, html, json, re, sys, unicodedata, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

SUPPORTED_EXTS = {".docx", ".pdf", ".md", ".txt", ".xlsx", ".html"}
CLEAN_PATTERN = re.compile(
    r'[\u200b-\u200d\ufeff\u2060\u202a-\u202e\u2066-\u2069\x00-\x08\x0b\x0c\x0e-\x1f\x7f]'
    r'|\(cid:\d+\)|\ufffd'
    r'|^\s*(?:[-—~]*\s*\d+\s*[-—~]*|Page\s+\d+(?:\s*(?:of|/)\s*\d+)?|\d+\s*/\s*\d+)\s*$',
    re.MULTILINE
)
MEETING_PATTERN = re.compile(
    r"^(?:#{1,4}[ \t]+|\[?[0-9]{1,2}:[0-9]{2}(?::[0-9]{2})?\]?[ \t]*|(?:議程|議題|Agenda|決議|待辦事項|Action Items|發言人)[0-9一二三四五六七八九十]*[:：]?[ \t]*)[^\r\n]*$",
    re.MULTILINE | re.IGNORECASE
)


def parse_html(raw: str) -> str:
    """Strip web clutter and elevate headings to Markdown anchors."""
    t = re.sub(r'<(script|style|nav|header|footer)[^>]*>.*?</\1>', '', raw, flags=re.DOTALL | re.IGNORECASE)
    t = re.sub(r'<h[1-6][^>]*>(.*?)</h[1-6]>', r'\n# \1\n', t, flags=re.DOTALL | re.IGNORECASE)
    return html.unescape(re.sub(r'<[^>]+>', ' ', t))


def read_raw(p: Path) -> str:
    """Format ingestion: docx, pdf, xlsx, html, md, txt."""
    ext = p.suffix.lower()
    if ext == ".docx":
        with zipfile.ZipFile(p) as z:
            tree = ET.fromstring(z.read("word/document.xml"))
            ns = {"w": "http" + chr(58) + chr(47) + chr(47) + "schemas.openxmlformats.org/wordprocessingml/2006/main"}
            return "\n".join("".join(t.text for t in n.findall(".//w:t", ns) if t.text) for n in tree.findall(".//w:p", ns))
    if ext == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
            lines = [f"# Sheet: {s}\n" + "\n".join(" | ".join(str(v).strip() for v in r if v is not None and str(v).strip()) for r in wb[s].iter_rows(values_only=True) if any(r)) for s in wb.sheetnames]
            return "\n".join(lines)
        except ImportError:
            sys.stderr.write("Warning: openpyxl not installed. Skipping XLSX.\n")
            return ""
    if ext == ".pdf":
        try:
            import pypdf
            return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(str(p)).pages)
        except Exception:
            return ""
    if ext in (".md", ".txt", ".html"):
        content = p.read_text(encoding="utf-8-sig", errors="replace")
        return parse_html(content) if ext == ".html" else content
    return ""


def clean_text(raw: str) -> str:
    """Single-pass sanitization: NFKC normalization + noise removal."""
    if not raw: return ""
    text = re.sub(r'(\b[A-Za-z]+)-\n([A-Za-z]+\b)', r'\1\2', unicodedata.normalize("NFKC", raw))
    return "\n".join(L.strip() for L in CLEAN_PATTERN.sub("", text).splitlines() if L.strip())


def chunk_content(text: str, source: str) -> list[dict]:
    """Segment content via meeting anchors or fallback to 500-line sliding window."""
    matches = list(MEETING_PATTERN.finditer(text))
    if matches:
        chunks = []
        for i, m in enumerate(matches):
            start = m.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = text[start:end].strip()
            if body:
                chunks.append({"source": source, "topic": m.group(0).strip("# ").strip(), "content": body})
        if chunks:
            return chunks

    lines = text.splitlines()
    total = len(lines)
    if total <= 500:
        return [{"source": source, "topic": "Full Document", "content": text}]
    chunks = []
    for idx, start_idx in enumerate(range(0, total, 450), 1):
        end_idx = min(start_idx + 500, total)
        chunks.append({
            "source": source, "topic": f"Part {idx:02d} (Lines {start_idx+1}-{end_idx})",
            "content": "\n".join(lines[start_idx:end_idx])
        })
        if end_idx >= total: break
    return chunks


def main():
    p = argparse.ArgumentParser(description="noai-note Ingester: Multi-source Noise Stripping & Meeting Chunking")
    p.add_argument("input", help="Target file or directory path")
    p.add_argument("-o", "--output", help="Destination file path (e.g. scratch/stream.txt)")
    p.add_argument("--format", choices=["stream", "json"], default="stream", help="Output serialization format")
    args = p.parse_args()

    inp = Path(args.input).resolve()
    files = [inp] if inp.is_file() and inp.suffix.lower() in SUPPORTED_EXTS else sorted(x for x in inp.rglob("*") if x.suffix.lower() in SUPPORTED_EXTS) if inp.is_dir() else []
    if not files:
        sys.exit(f"Error: No valid documents (.docx, .pdf, .xlsx, .html, .md, .txt) found in: {inp}")

    records = [c for f in files if (raw := read_raw(f)) for c in chunk_content(clean_text(raw), f.name) if c["content"]]
    if not records:
        sys.exit("Error: No extractable content found.")

    est_tokens = sum(len(r["content"]) for r in records) // 3
    out = json.dumps(records, ensure_ascii=False, indent=2) if args.format == "json" else "\n\n".join(
        f"=== [SOURCE: {r['source']} | TOPIC: {r['topic']}] ===\n{r['content']}" for r in records
    )

    if args.output:
        dest = Path(args.output).resolve()
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_text(out, encoding="utf-8-sig")
        sys.stderr.write(f"[OK] Ingested {len(files)} source(s), {len(records)} section(s) (約 {est_tokens:,} tokens) -> {dest.name}\n")
    else:
        print(out)


if __name__ == "__main__":
    main()
